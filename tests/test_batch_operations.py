import shutil
import sqlite3
import sys
import tempfile
import types
import unittest
from pathlib import Path

from openpyxl import load_workbook

try:
    import ollama  # noqa: F401
except (ModuleNotFoundError, ImportError):
    ollama_stub = types.ModuleType("ollama")

    class AsyncClientStub:
        def __init__(self, *args, **kwargs) -> None:
            pass

    ollama_stub.AsyncClient = AsyncClientStub
    sys.modules["ollama"] = ollama_stub

from app.batch_operations import execute_excel_batch, execute_sqlite_batch
from app.excel_operations import build_operation_preview, execute_operation
from app.excel_planner import ColumnRef, ExcelOperationPlan, FilterCondition, FilterOperator
from app.excel_service import inspect_workbook
from app.intent import Action, IntentBatch, IntentDraft, SourceType
from app.sql_operations import build_sql_preview, execute_sql_operation
from app.sql_planner import SqlAssignment, SqlColumnRef, SqlFilterCondition, SqlOperationPlan
from app.sql_service import inspect_sqlite


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEMO_FILE = PROJECT_ROOT / "workspace_files" / "price.xlsx"
if not DEMO_FILE.is_file():
    DEMO_FILE = PROJECT_ROOT / "demo_data" / "price.xlsx"


def _excel_column(metadata, header: str) -> ColumnRef:
    sheet = metadata.sheets[0]
    column = next(item for item in sheet.columns if item.header == header)
    return ColumnRef(index=column.index, header=column.header)


def _create_sqlite(path: Path) -> None:
    connection = sqlite3.connect(path)
    try:
        connection.execute(
            "CREATE TABLE products (id INTEGER PRIMARY KEY, category TEXT, stock INTEGER)"
        )
        connection.executemany(
            "INSERT INTO products(category, stock) VALUES (?, ?)",
            [("Периферия", 1), ("Периферия", 2), ("Аудио", 0)],
        )
        connection.commit()
    finally:
        connection.close()


def _sql_column(metadata, name: str) -> SqlColumnRef:
    table = metadata.tables[0]
    column = next(item for item in table.columns if item.name == name)
    return SqlColumnRef(name=column.name, declared_type=column.declared_type)


class BatchOperationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_excel_batch_uses_result_of_previous_task(self) -> None:
        source = self.root / "price.xlsx"
        planning = self.root / "planning.xlsx"
        shutil.copy2(DEMO_FILE, source)
        shutil.copy2(DEMO_FILE, planning)

        metadata1 = inspect_workbook(planning, planning.name)
        rename = ExcelOperationPlan(
            action=Action.RENAME_COLUMN,
            resolved=True,
            sheet_name=metadata1.sheets[0].name,
            target_columns=[_excel_column(metadata1, "Товар")],
            new_column_name="Название",
            confidence=0.99,
            resolution_note="test",
        )
        preview1 = build_operation_preview(planning, rename, metadata1)
        execute_operation(
            planning,
            planning.name,
            "planning-1",
            rename,
            preview1,
            metadata1,
            self.root / "planning-snapshots",
        )

        metadata2 = inspect_workbook(planning, planning.name)
        drop_renamed = ExcelOperationPlan(
            action=Action.DROP_COLUMN,
            resolved=True,
            sheet_name=metadata2.sheets[0].name,
            target_columns=[_excel_column(metadata2, "Название")],
            confidence=0.99,
            resolution_note="test",
        )
        preview2 = build_operation_preview(planning, drop_renamed, metadata2)

        outcome = execute_excel_batch(
            source,
            source.name,
            "batch-excel",
            [(rename, preview1, metadata1), (drop_renamed, preview2, metadata2)],
            self.root / "snapshots",
        )

        self.assertTrue(outcome.result_verified)
        self.assertTrue(outcome.snapshot_path.is_file())
        workbook = load_workbook(source)
        try:
            sheet = workbook[metadata1.sheets[0].name]
            headers = [
                sheet.cell(metadata1.sheets[0].header_row, index).value
                for index in range(1, (sheet.max_column or 0) + 1)
            ]
            self.assertNotIn("Товар", headers)
            self.assertNotIn("Название", headers)
        finally:
            workbook.close()

    def test_sqlite_batch_plans_second_task_after_first(self) -> None:
        source = self.root / "warehouse.sqlite3"
        planning = self.root / "planning.sqlite3"
        _create_sqlite(source)
        shutil.copy2(source, planning)

        metadata1 = inspect_sqlite(planning, planning.name)
        update = SqlOperationPlan(
            action=Action.UPDATE_ROWS,
            resolved=True,
            table_name="products",
            filters=[
                SqlFilterCondition(
                    column=_sql_column(metadata1, "category"),
                    operator=FilterOperator.EQ,
                    value="Периферия",
                )
            ],
            assignments=[SqlAssignment(column=_sql_column(metadata1, "stock"), value=77)],
            confidence=0.99,
            resolution_note="test",
        )
        preview1 = build_sql_preview(planning, update, metadata1)
        execute_sql_operation(
            planning,
            "planning-1",
            update,
            preview1,
            metadata1,
            self.root / "planning-snapshots",
        )

        metadata2 = inspect_sqlite(planning, planning.name)
        delete = SqlOperationPlan(
            action=Action.DELETE_ROWS,
            resolved=True,
            table_name="products",
            filters=[
                SqlFilterCondition(
                    column=_sql_column(metadata2, "stock"),
                    operator=FilterOperator.EQ,
                    value=77,
                )
            ],
            confidence=0.99,
            resolution_note="test",
        )
        preview2 = build_sql_preview(planning, delete, metadata2)
        self.assertEqual(preview2.matched_rows, 2)

        outcome = execute_sqlite_batch(
            source,
            source.name,
            "batch-sql",
            [(update, preview1, metadata1), (delete, preview2, metadata2)],
            self.root / "snapshots",
        )
        self.assertTrue(outcome.result_verified)
        connection = sqlite3.connect(source)
        try:
            rows = connection.execute(
                "SELECT category, stock FROM products ORDER BY id"
            ).fetchall()
            self.assertEqual(rows, [("Аудио", 0)])
        finally:
            connection.close()

    def test_intent_batch_propagates_shared_file_context(self) -> None:
        def task(action: Action, source: str | None) -> IntentDraft:
            return IntentDraft(
                normalized_request=action.value,
                action=action,
                source_type_hint=SourceType.EXCEL if source else SourceType.UNKNOWN,
                source_name_hint=source,
                resource_name_hint="Прайс" if source else None,
                column_hints=[],
                filter_hints=[],
                value_hints=[],
                is_write_operation=action != Action.SELECT,
                is_destructive=False,
                needs_discovery=False,
                confidence=0.9,
                interpretation_note="test",
                clarification_question=None,
            )

        batch = IntentBatch(
            tasks=[
                task(Action.RENAME_COLUMN, None),
                task(Action.UPDATE_ROWS, "price.xlsx"),
            ]
        )
        self.assertEqual(batch.tasks[0].source_name_hint, "price.xlsx")
        self.assertEqual(batch.tasks[0].resource_name_hint, "Прайс")
        self.assertEqual(batch.tasks[0].source_type_hint, SourceType.EXCEL)


if __name__ == "__main__":
    unittest.main(verbosity=2)
