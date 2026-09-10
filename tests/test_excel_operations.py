import hashlib
import shutil
import sqlite3
import sys
import tempfile
import types
import unittest
from pathlib import Path

try:
    import ollama  # noqa: F401
except (ModuleNotFoundError, ImportError):
    ollama_stub = types.ModuleType("ollama")
    ollama_stub.AsyncClient = object
    sys.modules["ollama"] = ollama_stub

from openpyxl import load_workbook
from pydantic import ValidationError

from app.excel_operations import build_operation_preview, execute_operation
from app.conversation import (
    ReplyKind,
    classify_reply,
    is_replace_all_command,
    replacement_value_from_text,
)
from app.database import DataRepository
from app.excel_planner import (
    ColumnAssignment,
    ColumnRef,
    ExcelOperationPlan,
    FilterCondition,
    FilterOperator,
    try_fast_excel_plan,
)
from app.excel_service import inspect_workbook
from app.intent import Action, IntentDraft, SourceType
from app.source_catalog import SourceCatalog, SourceOrigin, match_sources
from app.sql_operations import build_sql_preview, execute_sql_operation
from app.sql_planner import (
    SqlAssignment,
    SqlColumnRef,
    SqlFilterCondition,
    SqlOperationPlan,
)
from app.sql_service import SqliteMetadata, inspect_sqlite


PROJECT_ROOT = Path(__file__).resolve().parents[1]
# Tests must use an immutable fixture. workspace_files/price.xlsx is a live file
# edited by the Telegram agent during manual QA, so using it here makes tests
# depend on previous user actions. Prefer demo_data and only fall back to the
# workspace copy for old checkouts that do not contain the demo fixture.
DEMO_FILE = PROJECT_ROOT / "demo_data" / "price.xlsx"
if not DEMO_FILE.is_file():
    DEMO_FILE = PROJECT_ROOT / "workspace_files" / "price.xlsx"


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def excel_plan(action: Action, **kwargs) -> ExcelOperationPlan:
    return ExcelOperationPlan(
        action=action,
        resolved=True,
        sheet_name="Прайс",
        confidence=0.99,
        resolution_note="test",
        **kwargs,
    )


def create_sqlite(path: Path) -> None:
    connection = sqlite3.connect(path)
    try:
        connection.execute(
            """
            CREATE TABLE products (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                category TEXT NOT NULL,
                stock INTEGER NOT NULL,
                price REAL NOT NULL
            )
            """
        )
        connection.executemany(
            "INSERT INTO products(name, category, stock, price) VALUES (?, ?, ?, ?)",
            [
                ("Клавиатура", "Периферия", 12, 3490),
                ("Мышь", "Периферия", 4, 1890),
                ("Монитор 24", "Мониторы", 7, 17990),
                ("Наушники", "Аудио", 0, 5290),
                ("Веб-камера", "Периферия", 9, 4190),
                ("Микрофон", "Аудио", 2, 8990),
            ],
        )
        connection.commit()
    finally:
        connection.close()


def sql_column(metadata: SqliteMetadata, name: str) -> SqlColumnRef:
    column = next(item for item in metadata.tables[0].columns if item.name == name)
    return SqlColumnRef(name=column.name, declared_type=column.declared_type)


class DataOpsV4Tests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.excel = self.root / "price.xlsx"
        shutil.copy2(DEMO_FILE, self.excel)
        self.database = self.root / "warehouse.sqlite3"
        create_sqlite(self.database)
        self.snapshots = self.root / "snapshots"

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def execute_excel(self, plan: ExcelOperationPlan, operation_id: str):
        metadata = inspect_workbook(self.excel, self.excel.name)
        preview = build_operation_preview(self.excel, plan, metadata)
        before = file_hash(self.excel)
        outcome = execute_operation(
            self.excel,
            self.excel.name,
            operation_id,
            plan,
            preview,
            metadata,
            self.snapshots / "excel",
        )
        self.assertTrue(outcome.source_updated)
        self.assertTrue(outcome.result_verified)
        self.assertEqual(outcome.result_path, self.excel.resolve())
        self.assertEqual(file_hash(outcome.snapshot_path), before)
        self.assertNotEqual(file_hash(self.excel), before)
        return preview

    def test_catalog_local_and_telegram_sources(self) -> None:
        workspace = self.root / "workspace"
        telegram = self.root / "telegram"
        workspace.mkdir()
        shutil.copy2(self.excel, workspace / "price.xlsx")
        catalog = SourceCatalog(workspace, telegram)
        catalog.initialize()
        upload = catalog.allocate_telegram_upload(42, "from_chat.xlsx")
        shutil.copy2(self.excel, upload)
        sources = catalog.list_sources(42)
        self.assertEqual(len(sources), 2)
        self.assertEqual(
            {item.origin for item in sources},
            {SourceOrigin.LOCAL, SourceOrigin.TELEGRAM},
        )
        self.assertEqual(
            [item.display_name for item in match_sources("прайс", sources)],
            ["price.xlsx"],
        )

    def test_excel_select_is_read_only(self) -> None:
        metadata = inspect_workbook(self.excel, self.excel.name)
        plan = excel_plan(
            Action.SELECT,
            filters=[
                FilterCondition(
                    column=ColumnRef(index=5, header="Остаток"),
                    operator=FilterOperator.LT,
                    value=5,
                )
            ],
        )
        before = file_hash(self.excel)
        preview = build_operation_preview(self.excel, plan, metadata)
        self.assertFalse(preview.is_write)
        self.assertEqual(preview.matched_rows, 3)
        self.assertEqual(file_hash(self.excel), before)

    def test_excel_rename_and_add_write_in_place(self) -> None:
        rename = excel_plan(
            Action.RENAME_COLUMN,
            target_columns=[ColumnRef(index=2, header="Товар")],
            new_column_name="Наименование",
        )
        self.execute_excel(rename, "rename")
        workbook = load_workbook(self.excel)
        try:
            self.assertEqual(workbook["Прайс"]["B3"].value, "Наименование")
            self.assertTrue(workbook["Прайс"]["B3"].font.bold)
        finally:
            workbook.close()

        metadata = inspect_workbook(self.excel, self.excel.name)
        add = ExcelOperationPlan(
            action=Action.ADD_COLUMN,
            resolved=True,
            sheet_name="Прайс",
            new_column_name="Статус",
            new_column_default="Активен",
            fill_new_column=True,
            confidence=0.99,
            resolution_note="test",
        )
        preview = build_operation_preview(self.excel, add, metadata)
        execute_operation(
            self.excel,
            self.excel.name,
            "add",
            add,
            preview,
            metadata,
            self.snapshots / "excel",
        )
        workbook = load_workbook(self.excel)
        try:
            sheet = workbook["Прайс"]
            self.assertEqual(sheet["G3"].value, "Статус")
            self.assertEqual(
                [sheet.cell(row, 7).value for row in range(4, 10)],
                ["Активен"] * 6,
            )
        finally:
            workbook.close()

    def test_excel_update_and_drop(self) -> None:
        update = excel_plan(
            Action.UPDATE_ROWS,
            filters=[
                FilterCondition(
                    column=ColumnRef(index=3, header="Категория"),
                    operator=FilterOperator.EQ,
                    value="Периферия",
                )
            ],
            assignments=[
                ColumnAssignment(
                    column=ColumnRef(index=5, header="Остаток"),
                    value=99,
                )
            ],
        )
        self.assertEqual(self.execute_excel(update, "update").matched_rows, 3)
        metadata = inspect_workbook(self.excel, self.excel.name)
        drop = ExcelOperationPlan(
            action=Action.DROP_COLUMN,
            resolved=True,
            sheet_name="Прайс",
            target_columns=[ColumnRef(index=4, header="Цена")],
            confidence=0.99,
            resolution_note="test",
        )
        preview = build_operation_preview(self.excel, drop, metadata)
        execute_operation(
            self.excel,
            self.excel.name,
            "drop",
            drop,
            preview,
            metadata,
            self.snapshots / "excel",
        )
        workbook = load_workbook(self.excel)
        try:
            headers = [workbook["Прайс"].cell(3, index).value for index in range(1, 6)]
            self.assertNotIn("Цена", headers)
        finally:
            workbook.close()

    def test_excel_rejects_changed_source(self) -> None:
        metadata = inspect_workbook(self.excel, self.excel.name)
        plan = excel_plan(
            Action.RENAME_COLUMN,
            target_columns=[ColumnRef(index=2, header="Товар")],
            new_column_name="Название",
        )
        preview = build_operation_preview(self.excel, plan, metadata)
        workbook = load_workbook(self.excel)
        try:
            workbook["Прайс"]["A4"] = "CHANGED"
            workbook.save(self.excel)
        finally:
            workbook.close()
        with self.assertRaisesRegex(ValueError, "изменился после preview"):
            execute_operation(
                self.excel,
                self.excel.name,
                "conflict",
                plan,
                preview,
                metadata,
                self.snapshots / "excel",
            )

    def test_replace_all_filled_cells_in_workbook(self) -> None:
        metadata = inspect_workbook(self.excel, self.excel.name)
        plan = ExcelOperationPlan(
            action=Action.REPLACE_ALL_VALUES,
            resolved=True,
            sheet_name=None,
            replacement_value="НАНАНА",
            confidence=0.99,
            resolution_note="test",
        )
        preview = self.execute_excel(plan, "replace-all")
        self.assertGreater(preview.affected_cells, 0)
        self.assertIn("заполненных ячеек", preview.summary)
        workbook = load_workbook(self.excel)
        try:
            non_empty = [
                cell.value
                for worksheet in workbook.worksheets
                for cell in getattr(worksheet, "_cells", {}).values()
                if cell.value is not None
            ]
            self.assertTrue(non_empty)
            self.assertEqual(set(non_empty), {"НАНАНА"})
            self.assertIsNone(workbook["Прайс"]["A2"].value)
        finally:
            workbook.close()

    def test_natural_confirmation_and_replace_all_detection(self) -> None:
        for text in (
            "да",
            "давай",
            "делай",
            "да оно",
            "да, давай, делай",
            "да сделай это",
            "подтверждаю",
            "ок",
        ):
            self.assertEqual(classify_reply(text), ReplyKind.CONFIRM)
        for text in ("нет", "отмена", "не надо", "стоп"):
            self.assertEqual(classify_reply(text), ReplyKind.CANCEL)
        self.assertEqual(
            classify_reply("у периферии поставь остаток 99"),
            ReplyKind.OTHER,
        )
        command = "измени все ячейки на НАНАНАНА, вообще все ячейки"
        self.assertTrue(is_replace_all_command(command))
        self.assertEqual(replacement_value_from_text(command), "НАНАНАНА")
        self.assertFalse(
            is_replace_all_command("измени все ячейки в столбце остаток на 99")
        )

    def test_fast_language_plan_is_disabled_in_ai_native_mode(self) -> None:
        metadata = inspect_workbook(self.excel, self.excel.name)
        intent = IntentDraft(
            normalized_request="Установить остаток 99 для категории Периферия",
            action=Action.UPDATE_ROWS,
            source_type_hint=SourceType.UNKNOWN,
            source_name_hint=None,
            resource_name_hint=None,
            column_hints=["Категория", "Остаток"],
            filter_hints=["Категория = Периферия"],
            value_hints=["99"],
            is_write_operation=True,
            is_destructive=False,
            needs_discovery=False,
            confidence=0.9,
            interpretation_note="test",
            clarification_question=None,
        )
        # v5.3 deliberately does not parse Russian in Python. The compatibility
        # hook stays importable, but all semantic planning goes through the LLM.
        self.assertIsNone(
            try_fast_excel_plan(
                "у товаров категории периферия поставь остаток 99",
                intent,
                metadata,
            )
        )

    def test_dashboard_state_and_only_latest_operation_can_confirm(self) -> None:
        audit_database = self.root / "audit.sqlite3"
        repository = DataRepository(audit_database)
        repository.initialize()
        repository.set_dashboard(42, 42, 777, 2)
        dashboard = repository.get_dashboard(42)
        self.assertIsNotNone(dashboard)
        self.assertEqual((dashboard.chat_id, dashboard.message_id, dashboard.page), (42, 777, 2))

        repository.create_request("planned", 42, "первая команда")
        repository.save_plan("planned", 42, {"engine": "excel"})
        self.assertEqual(
            repository.get_latest_confirmable_operation(42).operation_id,
            "planned",
        )
        repository.create_request("newer", 42, "неполная новая команда")
        repository.stop_operation(
            "newer",
            42,
            "needs_clarification",
            "Не хватает значения",
            "clarification_required",
        )
        self.assertIsNone(repository.get_latest_confirmable_operation(42))

    def test_repository_connection_is_closed_after_use(self) -> None:
        repository = DataRepository(self.root / "connection-check.sqlite3")
        with repository._connect() as connection:
            self.assertEqual(connection.execute("SELECT 1").fetchone()[0], 1)

        with self.assertRaises(sqlite3.ProgrammingError):
            connection.execute("SELECT 1")

    def test_sqlite_select_and_update_transaction(self) -> None:
        metadata = inspect_sqlite(self.database, self.database.name)
        select = SqlOperationPlan(
            action=Action.SELECT,
            resolved=True,
            table_name="products",
            filters=[
                SqlFilterCondition(
                    column=sql_column(metadata, "stock"),
                    operator=FilterOperator.LT,
                    value=5,
                )
            ],
            confidence=0.99,
            resolution_note="test",
        )
        self.assertEqual(build_sql_preview(self.database, select, metadata).matched_rows, 3)

        update = SqlOperationPlan(
            action=Action.UPDATE_ROWS,
            resolved=True,
            table_name="products",
            filters=[
                SqlFilterCondition(
                    column=sql_column(metadata, "category"),
                    operator=FilterOperator.EQ,
                    value="Периферия",
                )
            ],
            assignments=[SqlAssignment(column=sql_column(metadata, "stock"), value=77)],
            confidence=0.99,
            resolution_note="test",
        )
        preview = build_sql_preview(self.database, update, metadata)
        outcome = execute_sql_operation(
            self.database,
            "sql-update",
            update,
            preview,
            metadata,
            self.snapshots / "sqlite",
        )
        self.assertTrue(outcome.transaction_committed)
        connection = sqlite3.connect(self.database)
        try:
            values = [
                row[0]
                for row in connection.execute(
                    "SELECT stock FROM products WHERE category = 'Периферия'"
                )
            ]
            self.assertEqual(values, [77, 77, 77])
        finally:
            connection.close()

    def test_sqlite_delete_and_filter_guard(self) -> None:
        metadata = inspect_sqlite(self.database, self.database.name)
        with self.assertRaises(ValidationError):
            SqlOperationPlan(
                action=Action.DELETE_ROWS,
                resolved=True,
                table_name="products",
                confidence=0.99,
                resolution_note="unsafe",
            )
        delete = SqlOperationPlan(
            action=Action.DELETE_ROWS,
            resolved=True,
            table_name="products",
            filters=[
                SqlFilterCondition(
                    column=sql_column(metadata, "stock"),
                    operator=FilterOperator.EQ,
                    value=0,
                )
            ],
            confidence=0.99,
            resolution_note="test",
        )
        preview = build_sql_preview(self.database, delete, metadata)
        self.assertEqual(preview.matched_rows, 1)
        execute_sql_operation(
            self.database,
            "sql-delete",
            delete,
            preview,
            metadata,
            self.snapshots / "sqlite",
        )
        connection = sqlite3.connect(self.database)
        try:
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM products WHERE stock = 0"
                ).fetchone()[0],
                0,
            )
        finally:
            connection.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
