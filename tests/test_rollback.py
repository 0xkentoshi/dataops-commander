from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.database import DataRepository
from app.rollback import restore_snapshot, sha256_file


class RollbackTests(unittest.TestCase):
    def _write_book(self, path: Path, value: str) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "value"
        ws["A2"] = value
        wb.save(path)
        wb.close()

    def test_restore_snapshot_requires_current_version_and_restores_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "data.xlsx"
            snapshot = root / "before.xlsx"
            self._write_book(source, "before")
            snapshot.write_bytes(source.read_bytes())
            self._write_book(source, "after")
            result_hash = sha256_file(source)

            outcome = restore_snapshot(
                source,
                snapshot,
                expected_current_sha256=result_hash,
                snapshots_root=root / "undo",
            )
            self.assertTrue(outcome.snapshot_path.is_file())
            wb = load_workbook(source)
            try:
                self.assertEqual(wb.active["A2"].value, "before")
            finally:
                wb.close()

    def test_restore_snapshot_rejects_external_change(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "data.xlsx"
            snapshot = root / "before.xlsx"
            self._write_book(source, "before")
            snapshot.write_bytes(source.read_bytes())
            self._write_book(source, "after")
            old_hash = sha256_file(source)
            self._write_book(source, "manual edit")

            with self.assertRaisesRegex(ValueError, "изменился"):
                restore_snapshot(
                    source,
                    snapshot,
                    expected_current_sha256=old_hash,
                    snapshots_root=root / "undo",
                )

    def test_repository_tracks_latest_undoable_and_marks_target_undone(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            db = DataRepository(root / "audit.sqlite3")
            db.initialize()
            source = root / "data.xlsx"
            snapshot = root / "snapshot.xlsx"
            self._write_book(source, "after")
            self._write_book(snapshot, "before")

            db.create_request("op1", 42, "change")
            db.save_intent("op1", 42, "update_rows", "excel", {})
            db.save_plan("op1", 42, {"action": "update_rows"})
            db.claim_operation("op1", 42)
            db.complete_operation("op1", 42, snapshot, source, sha256_file(source))

            latest = db.get_latest_undoable_operation(42, source)
            self.assertIsNotNone(latest)
            self.assertEqual(latest.operation_id, "op1")
            self.assertTrue(db.mark_undone("op1", 42, "undo1"))
            self.assertIsNone(db.get_latest_undoable_operation(42, source))


if __name__ == "__main__":
    unittest.main()
