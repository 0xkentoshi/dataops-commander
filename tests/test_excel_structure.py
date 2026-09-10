from __future__ import annotations

import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from app.excel_service import inspect_workbook


class ExcelStructureTests(unittest.TestCase):
    @staticmethod
    def _corrupt_dimension(source: Path, target: Path, ref: str = "A1:A1") -> None:
        with zipfile.ZipFile(source, "r") as reader, zipfile.ZipFile(
            target, "w", zipfile.ZIP_DEFLATED
        ) as writer:
            for item in reader.infolist():
                data = reader.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    text = data.decode("utf-8")
                    text = re.sub(
                        r'<dimension ref="[^"]+"',
                        f'<dimension ref="{ref}"',
                        text,
                        count=1,
                    )
                    data = text.encode("utf-8")
                writer.writestr(item, data)

    def test_bad_xlsx_dimension_does_not_hide_real_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            normal = root / "clients.xlsx"
            broken = root / "clients_bad_dimension.xlsx"

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "clients"
            sheet.append(["client_id", "client", "email", "city"])
            for index in range(1, 1001):
                sheet.append(
                    [
                        index,
                        f"Client {index}",
                        f"user{index}@example.com",
                        "Новосибирск" if index % 2 == 0 else "Москва",
                    ]
                )
            workbook.save(normal)
            workbook.close()
            self._corrupt_dimension(normal, broken)

            metadata = inspect_workbook(broken, broken.name)
            clients = metadata.sheets[0]

            self.assertEqual(clients.data_rows, 1000)
            self.assertEqual(clients.detected_columns, 4)
            self.assertEqual(
                [column.header for column in clients.columns],
                ["client_id", "client", "email", "city"],
            )

    def test_header_can_be_below_first_ten_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "clients"
            sheet["A1"] = "Отчёт по клиентам"
            headers = ["client_id", "client", "email", "city"]
            for column, value in enumerate(headers, start=1):
                sheet.cell(15, column).value = value
            for index in range(1, 11):
                row = 15 + index
                sheet.cell(row, 1).value = index
                sheet.cell(row, 2).value = f"C{index}"
                sheet.cell(row, 3).value = f"u{index}@x.ru"
                sheet.cell(row, 4).value = "Новосибирск"
            workbook.save(path)
            workbook.close()

            metadata = inspect_workbook(path, path.name)
            clients = metadata.sheets[0]
            self.assertEqual(clients.header_row, 15)
            self.assertEqual(clients.data_rows, 10)
            self.assertEqual(clients.detected_columns, 4)


if __name__ == "__main__":
    unittest.main()
