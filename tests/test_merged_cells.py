import tempfile
import unittest
import sys
import types
from pathlib import Path

try:
    import ollama  # noqa: F401
except (ModuleNotFoundError, ImportError):
    ollama_stub = types.ModuleType("ollama")
    ollama_stub.AsyncClient = object
    sys.modules["ollama"] = ollama_stub

from openpyxl import Workbook, load_workbook

from app.excel_operations import build_operation_preview, execute_operation
from app.excel_planner import ColumnAssignment, ColumnRef, ExcelOperationPlan
from app.excel_service import inspect_workbook
from app.intent import Action


class MergedCellWriteTests(unittest.TestCase):
    def _book(self, root: Path) -> Path:
        path = root / "merged.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Данные"
        ws.append(["Артикул", "Товар", "Категория"])
        ws.append(["1", "Один", "Аудио"])
        ws.append(["2", "Два", "Периферия"])
        ws.append(["3", "Три", None])
        ws.merge_cells("C3:C4")
        wb.save(path)
        wb.close()
        return path

    def test_clear_values_skips_read_only_merged_proxy(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = self._book(root)
            metadata = inspect_workbook(path, path.name)
            plan = ExcelOperationPlan(
                action=Action.CLEAR_VALUES,
                resolved=True,
                sheet_name="Данные",
                confidence=1.0,
                resolution_note="test",
                target_columns=[ColumnRef(index=3, header="Категория")],
            )
            preview = build_operation_preview(path, plan, metadata)
            self.assertGreater(preview.affected_cells, 0)
            outcome = execute_operation(
                path,
                path.name,
                "merged-clear",
                plan,
                preview,
                metadata,
                snapshots_root=root / "snapshots",
            )
            self.assertTrue(outcome.result_verified)
            wb = load_workbook(path)
            try:
                ws = wb["Данные"]
                self.assertIsNone(ws["C2"].value)
                self.assertIsNone(ws["C3"].value)
                self.assertIsNone(ws["C4"].value)
                self.assertIn("C3:C4", {str(r) for r in ws.merged_cells.ranges})
            finally:
                wb.close()

    def test_update_into_merged_proxy_is_rejected_before_confirm(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = self._book(root)
            metadata = inspect_workbook(path, path.name)
            plan = ExcelOperationPlan(
                action=Action.UPDATE_ROWS,
                resolved=True,
                sheet_name="Данные",
                confidence=1.0,
                resolution_note="test",
                apply_to_all_rows=True,
                assignments=[
                    ColumnAssignment(
                        column=ColumnRef(index=3, header="Категория"),
                        value="Новое",
                    )
                ],
            )
            with self.assertRaisesRegex(ValueError, "объединённый диапазон"):
                build_operation_preview(path, plan, metadata)


if __name__ == "__main__":
    unittest.main()
