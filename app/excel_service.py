from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field


MAX_COLUMNS_TO_SCAN = 200
MAX_DATA_ROWS_TO_SCAN = 10_000
MAX_SAMPLE_VALUES = 5
MAX_SAMPLE_LENGTH = 80
HEADER_PROBE_ROWS = 50


class ColumnMetadata(BaseModel):
    model_config = ConfigDict(extra="forbid")

    index: int
    letter: str
    header: str
    non_empty_cells: int
    formula_cells: int
    samples: list[str] = Field(default_factory=list)


class SheetMetadata(BaseModel):
    model_config = ConfigDict(extra="forbid")

    name: str
    header_row: int
    data_rows: int
    detected_columns: int
    columns: list[ColumnMetadata] = Field(default_factory=list)


class WorkbookMetadata(BaseModel):
    model_config = ConfigDict(extra="forbid")

    original_name: str
    sheets: list[SheetMetadata] = Field(default_factory=list)
    total_formula_cells: int = 0
    warnings: list[str] = Field(default_factory=list)


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _stringify(value: Any) -> str:
    if isinstance(value, (datetime, date, time)):
        text = value.isoformat()
    else:
        text = str(value)
    text = text.replace("\n", " ").strip()
    return text if len(text) <= MAX_SAMPLE_LENGTH else f"{text[:79]}…"


def _detect_header_row_from_rows(rows: list[tuple[Any, ...]]) -> int:
    """Определяет строку заголовков по верхней части листа.

    Широкая текстовая строка перед строками похожей ширины предпочтительнее
    одиночных заголовков/титулов. Это важно для реальных выгрузок, где перед
    таблицей могут быть название отчёта, дата, служебные строки и пустоты.
    """
    best_row = 1
    best_score = float("-inf")
    for offset, row in enumerate(rows, start=1):
        values = [cell.value for cell in row]
        non_empty = [value for value in values if not _is_empty(value)]
        if not non_empty:
            continue
        width = len(non_empty)
        text_cells = sum(
            isinstance(value, str) and not value.startswith("=") for value in non_empty
        )
        unique_cells = len({_stringify(value).casefold() for value in non_empty})
        next_row_non_empty = 0
        if offset < len(rows):
            next_row_non_empty = sum(
                not _is_empty(cell.value) for cell in rows[offset]
            )
        shape_match = (
            min(width, next_row_non_empty) / max(width, next_row_non_empty)
            if next_row_non_empty
            else 0.0
        )
        score = (
            text_cells * 3
            + width * 1.5
            + unique_cells * 0.5
            + shape_match * 4
            - offset * 0.05
        )
        if width == 1 and next_row_non_empty > 1:
            score -= 4
        if score > best_score:
            best_score = score
            best_row = offset
    return best_row


def _repair_read_only_dimensions(worksheet: Any) -> bool:
    """Пересчитывает подозрительные XLSX dimensions.

    Некоторые экспортёры сохраняют в sheet XML неверный `<dimension ref>`,
    например A1:A1, хотя в книге реально тысячи строк и несколько столбцов.
    openpyxl в read_only-режиме доверяет этому ref и тогда видит только A1.
    `reset_dimensions()` + `calculate_dimension(force=True)` заставляет один
    раз пройти XML и восстановить фактический используемый диапазон.
    """
    max_row = worksheet.max_row or 0
    max_column = worksheet.max_column or 0
    if max_row > 1 and max_column > 1:
        return False
    reset = getattr(worksheet, "reset_dimensions", None)
    if not callable(reset):
        return False
    reset()
    worksheet.calculate_dimension(force=True)
    return True


def inspect_workbook(file_path: Path, original_name: str) -> WorkbookMetadata:
    """Линейно сканирует структуру XLSX.

    Старый вариант вызывал worksheet.cell() тысячи раз в read_only-режиме.
    Для больших XLSX это приводило к фактически квадратичному чтению XML и
    файлы на несколько тысяч строк могли выглядеть как зависшие. Здесь строки
    читаются последовательными iter_rows(), поэтому 3–10 тыс. строк — обычный
    рабочий размер для структурного анализа.
    """
    if not file_path.is_file():
        raise ValueError("Excel-файл больше не существует.")

    workbook = load_workbook(file_path, read_only=True, data_only=False)
    sheets: list[SheetMetadata] = []
    warnings: list[str] = []
    total_formula_cells = 0
    try:
        for worksheet in workbook.worksheets:
            repaired_dimensions = _repair_read_only_dimensions(worksheet)
            if repaired_dimensions:
                warnings.append(
                    f"Лист «{worksheet.title}»: фактический диапазон ячеек "
                    "пересчитан из XLSX, потому что служебные dimensions были подозрительными."
                )

            actual_max_column = max(worksheet.max_column or 1, 1)
            max_column = min(actual_max_column, MAX_COLUMNS_TO_SCAN)
            if actual_max_column > MAX_COLUMNS_TO_SCAN:
                warnings.append(
                    f"Лист «{worksheet.title}»: просмотрены только первые "
                    f"{MAX_COLUMNS_TO_SCAN} столбцов."
                )

            worksheet_max_row = max(worksheet.max_row or 1, 1)
            header_probe_end = min(worksheet_max_row, HEADER_PROBE_ROWS)
            first_rows = list(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=header_probe_end,
                    min_col=1,
                    max_col=max_column,
                    values_only=False,
                )
            )
            header_row = _detect_header_row_from_rows(first_rows)
            header_cells = first_rows[header_row - 1] if first_rows else tuple()

            last_row = min(
                worksheet_max_row,
                header_row + MAX_DATA_ROWS_TO_SCAN,
            )
            if worksheet_max_row > last_row:
                warnings.append(
                    f"Лист «{worksheet.title}»: статистика рассчитана по первым "
                    f"{MAX_DATA_ROWS_TO_SCAN} строкам данных."
                )

            stats = [
                {"non_empty": 0, "formula": 0, "samples": []}
                for _ in range(max_column)
            ]
            non_empty_data_rows = 0

            if last_row > header_row:
                for row in worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=last_row,
                    min_col=1,
                    max_col=max_column,
                    values_only=False,
                ):
                    row_has_data = False
                    for offset, cell in enumerate(row):
                        value = cell.value
                        if _is_empty(value):
                            continue
                        row_has_data = True
                        item = stats[offset]
                        item["non_empty"] += 1
                        if cell.data_type == "f" or (
                            isinstance(value, str) and value.startswith("=")
                        ):
                            item["formula"] += 1
                        sample = _stringify(value)
                        samples = item["samples"]
                        if len(samples) < MAX_SAMPLE_VALUES and sample not in samples:
                            samples.append(sample)
                    if row_has_data:
                        non_empty_data_rows += 1

            columns: list[ColumnMetadata] = []
            for column_index in range(1, max_column + 1):
                letter = get_column_letter(column_index)
                header_value = (
                    header_cells[column_index - 1].value
                    if column_index - 1 < len(header_cells)
                    else None
                )
                item = stats[column_index - 1]
                if _is_empty(header_value) and item["non_empty"] == 0:
                    continue
                header = (
                    _stringify(header_value)
                    if not _is_empty(header_value)
                    else f"<без названия {letter}>"
                )
                total_formula_cells += int(item["formula"])
                columns.append(
                    ColumnMetadata(
                        index=column_index,
                        letter=letter,
                        header=header,
                        non_empty_cells=int(item["non_empty"]),
                        formula_cells=int(item["formula"]),
                        samples=list(item["samples"]),
                    )
                )

            sheets.append(
                SheetMetadata(
                    name=worksheet.title,
                    header_row=header_row,
                    data_rows=non_empty_data_rows,
                    detected_columns=len(columns),
                    columns=columns,
                )
            )
    finally:
        workbook.close()

    if not sheets:
        warnings.append("В книге не найдено ни одного листа.")
    return WorkbookMetadata(
        original_name=original_name,
        sheets=sheets,
        total_formula_cells=total_formula_cells,
        warnings=warnings,
    )
