import hashlib
import json
import math
import os
import re
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, range_boundaries
from pydantic import BaseModel, ConfigDict, Field

from app.excel_planner import (
    CellValue,
    ColumnAssignment,
    ColumnRef,
    ExcelOperationPlan,
    FilterCondition,
    FilterOperator,
)
from app.excel_service import SheetMetadata, WorkbookMetadata
from app.intent import Action


MAX_PREVIEW_ROWS = 10
MAX_PREVIEW_COLUMNS = 8
MAX_MUTATED_ROWS = 5_000
MAX_MUTATED_CELLS = 50_000

WRITE_ACTIONS = {
    Action.ADD_COLUMN,
    Action.RENAME_COLUMN,
    Action.DROP_COLUMN,
    Action.UPDATE_ROWS,
    Action.REPLACE_ALL_VALUES,
    Action.DELETE_ROWS,
    Action.CLEAR_VALUES,
    Action.DEDUPLICATE,
}

STRUCTURAL_ACTIONS = {
    Action.DROP_COLUMN,
    Action.DELETE_ROWS,
    Action.DEDUPLICATE,
}


class PreviewCell(BaseModel):
    model_config = ConfigDict(extra="forbid")

    column_header: str
    before: str
    after: str | None = None


class PreviewRow(BaseModel):
    model_config = ConfigDict(extra="forbid")

    row_number: int
    sheet_name: str | None = None
    cells: list[PreviewCell] = Field(default_factory=list)


class ExcelOperationPreview(BaseModel):
    model_config = ConfigDict(extra="forbid")

    action: Action
    sheet_name: str
    is_write: bool
    has_changes: bool
    matched_rows: int
    affected_cells: int
    matched_row_numbers: list[int] = Field(default_factory=list)
    matched_cell_addresses: list[str] = Field(default_factory=list)
    rows: list[PreviewRow] = Field(default_factory=list)
    summary: str
    source_sha256: str


@dataclass(frozen=True, slots=True)
class ExcelExecutionOutcome:
    snapshot_path: Path
    result_path: Path
    changed_rows: int
    affected_cells: int
    columns_before: int
    columns_after: int
    source_updated: bool
    result_verified: bool


def _sha256(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _display(value: Any) -> str:
    if value is None:
        return "<empty>"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value).replace("\n", " ").strip()
    return text if len(text) <= 100 else f"{text[:99]}…"


def _number(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if not isinstance(value, str):
        return None
    cleaned = value.strip().replace(" ", "").replace(",", ".")
    try:
        number = float(cleaned)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    text = value.strip()
    for parser in (
        lambda item: datetime.fromisoformat(item).date(),
        lambda item: datetime.strptime(item, "%d.%m.%Y").date(),
        lambda item: datetime.strptime(item, "%d/%m/%Y").date(),
    ):
        try:
            return parser(text)
        except ValueError:
            continue
    return None


def _normalized(value: Any) -> Any:
    if _empty(value):
        return None
    if isinstance(value, str):
        return " ".join(value.casefold().split())
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return "NaN"
    return value


def _ordered_pair(actual: Any, expected: Any) -> tuple[Any, Any] | None:
    actual_number = _number(actual)
    expected_number = _number(expected)
    if actual_number is not None and expected_number is not None:
        return actual_number, expected_number
    actual_date = _date_value(actual)
    expected_date = _date_value(expected)
    if actual_date is not None and expected_date is not None:
        return actual_date, expected_date
    if isinstance(actual, str) and isinstance(expected, str):
        return actual.casefold().strip(), expected.casefold().strip()
    return None


def _matches(value: Any, condition: FilterCondition) -> bool:
    operator = condition.operator
    expected = condition.value
    if operator == FilterOperator.IS_EMPTY:
        return _empty(value)
    if operator == FilterOperator.NOT_EMPTY:
        return not _empty(value)
    if operator == FilterOperator.EQ:
        pair = _ordered_pair(value, expected)
        return pair[0] == pair[1] if pair else _normalized(value) == _normalized(expected)
    if operator == FilterOperator.NE:
        pair = _ordered_pair(value, expected)
        return pair[0] != pair[1] if pair else _normalized(value) != _normalized(expected)
    if operator in {FilterOperator.CONTAINS, FilterOperator.NOT_CONTAINS}:
        contained = str(expected).casefold().strip() in str(value or "").casefold()
        return contained if operator == FilterOperator.CONTAINS else not contained
    if operator == FilterOperator.OLDER_THAN_DAYS:
        actual_date = _date_value(value)
        days = _number(expected)
        if actual_date is None or days is None:
            return False
        threshold = datetime.now(timezone.utc).date() - timedelta(days=float(days))
        return actual_date < threshold
    pair = _ordered_pair(value, expected)
    if pair is None:
        return False
    left, right = pair
    if operator == FilterOperator.GT:
        return left > right
    if operator == FilterOperator.GTE:
        return left >= right
    if operator == FilterOperator.LT:
        return left < right
    if operator == FilterOperator.LTE:
        return left <= right
    return False


def _coerce_write(value: CellValue, current: Any) -> CellValue | date | datetime:
    if value is None or not isinstance(value, str):
        return value
    if isinstance(current, bool):
        normalized = value.strip().casefold()
        if normalized in {"true", "да", "yes", "1"}:
            return True
        if normalized in {"false", "нет", "no", "0"}:
            return False
    if isinstance(current, int) and not isinstance(current, bool):
        converted = _number(value)
        return int(converted) if converted is not None else value
    if isinstance(current, float):
        converted = _number(value)
        return float(converted) if converted is not None else value
    if isinstance(current, datetime):
        parsed = _date_value(value)
        return datetime.combine(parsed, current.time()) if parsed else value
    if isinstance(current, date):
        return _date_value(value) or value
    return value


def _sheet_metadata(metadata: WorkbookMetadata, sheet_name: str) -> SheetMetadata:
    result = next((sheet for sheet in metadata.sheets if sheet.name == sheet_name), None)
    if result is None:
        raise ValueError(f"Sheet “{sheet_name}” is not present in the schema.")
    return result


def _data_row_numbers(worksheet: Any, sheet: SheetMetadata) -> list[int]:
    indexes = [column.index for column in sheet.columns]
    result: list[int] = []
    for row_number in range(sheet.header_row + 1, (worksheet.max_row or 0) + 1):
        if any(not _empty(worksheet.cell(row_number, index).value) for index in indexes):
            result.append(row_number)
    return result


def _filtered_rows(
    worksheet: Any,
    row_numbers: list[int],
    filters: list[FilterCondition],
) -> list[int]:
    if not filters:
        return list(row_numbers)
    return [
        row_number
        for row_number in row_numbers
        if all(
            _matches(
                worksheet.cell(row_number, condition.column.index).value,
                condition,
            )
            for condition in filters
        )
    ]


def _duplicate_rows(
    worksheet: Any,
    row_numbers: list[int],
    keys: list[ColumnRef],
) -> list[int]:
    seen: set[tuple[Any, ...]] = set()
    duplicates: list[int] = []
    for row_number in row_numbers:
        key = tuple(
            _normalized(worksheet.cell(row_number, column.index).value)
            for column in keys
        )
        if key in seen:
            duplicates.append(row_number)
        else:
            seen.add(key)
    return duplicates


def _preview_columns(plan: ExcelOperationPlan, sheet: SheetMetadata) -> list[ColumnRef]:
    if plan.action == Action.SELECT:
        refs = plan.selected_columns or [
            ColumnRef(index=column.index, header=column.header) for column in sheet.columns
        ]
    elif plan.action == Action.UPDATE_ROWS:
        refs = [
            *(condition.column for condition in plan.filters),
            *(assignment.column for assignment in plan.assignments),
        ]
    elif plan.action == Action.DELETE_ROWS:
        refs = [ColumnRef(index=column.index, header=column.header) for column in sheet.columns]
    elif plan.action == Action.CLEAR_VALUES:
        refs = [
            *(condition.column for condition in plan.filters),
            *plan.target_columns,
        ]
    elif plan.action == Action.DEDUPLICATE:
        refs = plan.deduplicate_columns
    else:
        refs = plan.target_columns
    unique: dict[int, ColumnRef] = {}
    for ref in refs:
        unique.setdefault(ref.index, ref)
    return list(unique.values())[:MAX_PREVIEW_COLUMNS]


def _preview_rows(
    worksheet: Any,
    row_numbers: list[int],
    refs: list[ColumnRef],
    plan: ExcelOperationPlan,
) -> list[PreviewRow]:
    assignments = {item.column.index: item.value for item in plan.assignments}
    targets = {item.index for item in plan.target_columns}
    rows: list[PreviewRow] = []
    for row_number in row_numbers[:MAX_PREVIEW_ROWS]:
        cells: list[PreviewCell] = []
        for ref in refs:
            value = worksheet.cell(row_number, ref.index).value
            after: str | None = None
            if plan.action == Action.UPDATE_ROWS and ref.index in assignments:
                after = _display(_coerce_write(assignments[ref.index], value))
            elif plan.action == Action.CLEAR_VALUES and ref.index in targets:
                after = "<empty>"
            cells.append(
                PreviewCell(
                    column_header=ref.header,
                    before=_display(value),
                    after=after,
                )
            )
        rows.append(PreviewRow(row_number=row_number, cells=cells))
    return rows


def _structural_guard(worksheet: Any, action: Action) -> None:
    if action not in STRUCTURAL_ACTIONS:
        return
    if worksheet.tables:
        raise ValueError(
            "The sheet contains a formal Excel Table. Structural changes are "
            "blocked until its table range can be updated safely."
        )
    if getattr(worksheet, "_charts", None) or getattr(worksheet, "_images", None):
        raise ValueError(
            "The sheet contains charts or images. Structural changes are "
            "blocked to avoid damaging their references."
        )


def _workbook_has_formulas(workbook: Any) -> bool:
    return any(
        cell.data_type == "f"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )


def _non_empty_workbook_cells(workbook: Any) -> list[tuple[Any, Any]]:
    result: list[tuple[Any, Any]] = []
    for worksheet in workbook.worksheets:
        cells = sorted(
            getattr(worksheet, "_cells", {}).values(),
            key=lambda cell: (cell.row, cell.column),
        )
        result.extend(
            (worksheet, cell)
            for cell in cells
            if not _empty(getattr(cell, "value", None))
        )
    return result


def _replace_all_preview(
    workbook: Any,
    replacement: CellValue,
) -> tuple[int, int, int, list[PreviewRow]]:
    targets = _non_empty_workbook_cells(workbook)
    if len(targets) > MAX_MUTATED_CELLS:
        raise ValueError(
            f"The workbook contains {len(targets)} filled cells; the operation limit is "
            f"{MAX_MUTATED_CELLS}."
        )

    affected = sum(
        cell.value != _coerce_write(replacement, cell.value)
        for _, cell in targets
    )
    row_keys = {(worksheet.title, cell.row) for worksheet, cell in targets}
    preview_rows: list[PreviewRow] = []
    for worksheet, cell in targets:
        existing = next(
            (
                row
                for row in preview_rows
                if row.sheet_name == worksheet.title and row.row_number == cell.row
            ),
            None,
        )
        if existing is None:
            if len(preview_rows) >= MAX_PREVIEW_ROWS:
                continue
            existing = PreviewRow(
                sheet_name=worksheet.title,
                row_number=cell.row,
            )
            preview_rows.append(existing)
        if len(existing.cells) >= MAX_PREVIEW_COLUMNS:
            continue
        existing.cells.append(
            PreviewCell(
                column_header=get_column_letter(cell.column),
                before=_display(cell.value),
                after=_display(_coerce_write(replacement, cell.value)),
            )
        )
    return len(targets), len(row_keys), affected, preview_rows


def _apply_replace_all(workbook: Any, replacement: CellValue) -> None:
    for _, cell in _non_empty_workbook_cells(workbook):
        cell.value = _coerce_write(replacement, cell.value)



def _matching_cell_addresses(worksheet: Any, value: CellValue) -> list[str]:
    """Ищет точные логические совпадения по всему используемому листу.

    Поиск включает строку заголовков: команда вида «удали ячейку, где
    написано Цена» должна уметь clear сам заголовок, а не удалить строку.
    """
    expected = _normalized(value)
    result: list[str] = []
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if _empty(cell.value):
                continue
            if _normalized(cell.value) == expected:
                result.append(cell.coordinate)
                if len(result) > MAX_MUTATED_CELLS:
                    raise ValueError(
                        f"More than {MAX_MUTATED_CELLS} matching cells were found — "
                        "narrow the condition."
                    )
    return result


def _matching_cells_preview(worksheet: Any, addresses: list[str]) -> list[PreviewRow]:
    rows: list[PreviewRow] = []
    for address in addresses[:MAX_PREVIEW_ROWS]:
        cell = worksheet[address]
        rows.append(
            PreviewRow(
                sheet_name=worksheet.title,
                row_number=cell.row,
                cells=[
                    PreviewCell(
                        column_header=cell.coordinate,
                        before=_display(cell.value),
                        after="<empty>",
                    )
                ],
            )
        )
    return rows

def build_operation_preview(
    file_path: Path,
    plan: ExcelOperationPlan,
    metadata: WorkbookMetadata,
) -> ExcelOperationPreview:
    if not plan.resolved or (
        plan.action != Action.REPLACE_ALL_VALUES and plan.sheet_name is None
    ):
        raise ValueError("Cannot build a preview for an unresolved plan.")
    is_write = plan.action in WRITE_ACTIONS
    if is_write and metadata.total_formula_cells > 0:
        raise ValueError(
            "The workbook contains formulas. Writes are blocked until formula "
            "dependency checks are available."
        )
    source_hash = _sha256(file_path)
    workbook = load_workbook(file_path, read_only=False, data_only=False)
    try:
        if is_write and _workbook_has_formulas(workbook):
            raise ValueError(
                "Formulas were detected in the workbook. Writes are blocked until "
                "dependency checks are available."
            )
        if plan.action == Action.REPLACE_ALL_VALUES:
            total, matched_rows, affected, preview_rows = _replace_all_preview(
                workbook,
                plan.replacement_value,
            )
            return ExcelOperationPreview(
                action=plan.action,
                sheet_name=f"Entire workbook ({len(workbook.worksheets)} sheets)",
                is_write=True,
                has_changes=affected > 0,
                matched_rows=matched_rows,
                affected_cells=affected,
                matched_row_numbers=[],
                rows=preview_rows,
                summary=(
                    f"Replace all {total} filled cells in the workbook "
                    f"with “{_display(plan.replacement_value)}”. "
                    f"Cells that will actually change: {affected}."
                ),
                source_sha256=source_hash,
            )

        sheet = _sheet_metadata(metadata, plan.sheet_name or "")
        worksheet = workbook[plan.sheet_name]
        if plan.action == Action.CLEAR_VALUES and plan.match_cell_value is not None:
            addresses = _matching_cell_addresses(worksheet, plan.match_cell_value)
            if not plan.match_all_cells and len(addresses) > 1:
                sample = ", ".join(addresses[:6])
                raise ValueError(
                    f"Found {len(addresses)} cells with value “{_display(plan.match_cell_value)}” "
                    f"({sample}). Narrow the location or explicitly request clearing all matching cells."
                )
            affected = len(addresses)
            matched_rows = sorted({worksheet[address].row for address in addresses})
            return ExcelOperationPreview(
                action=plan.action,
                sheet_name=plan.sheet_name,
                is_write=True,
                has_changes=affected > 0,
                matched_rows=len(matched_rows),
                affected_cells=affected,
                matched_row_numbers=matched_rows,
                matched_cell_addresses=addresses,
                rows=_matching_cells_preview(worksheet, addresses),
                summary=(
                    f"Clear cells with value “{_display(plan.match_cell_value)}”: "
                    f"{affected}. Rows and columns will be preserved."
                ),
                source_sha256=source_hash,
            )
        _structural_guard(worksheet, plan.action)
        data_rows = _data_row_numbers(worksheet, sheet)
        matched = _filtered_rows(worksheet, data_rows, plan.filters)
        summary: str

        if plan.action == Action.DEDUPLICATE:
            matched = _duplicate_rows(
                worksheet,
                matched,
                plan.deduplicate_columns,
            )
            summary = f"Duplicate rows to remove: {len(matched)}."
        elif plan.action == Action.SELECT:
            summary = f"Rows found: {len(matched)}. The file will not be modified."
        elif plan.action == Action.RENAME_COLUMN:
            target = plan.target_columns[0]
            summary = f"Rename “{target.header}” to “{plan.new_column_name}»."
        elif plan.action == Action.ADD_COLUMN:
            fill_note = (
                f" and fill with “{_display(plan.new_column_default)}»"
                if plan.fill_new_column
                else " without filling"
            )
            summary = f"Add column “{plan.new_column_name}»{fill_note}."
        elif plan.action == Action.DROP_COLUMN:
            target = plan.target_columns[0]
            if len(sheet.columns) <= 1:
                raise ValueError("Cannot drop the only column in the table.")
            summary = f"Drop column “{target.header}»."
        elif plan.action == Action.UPDATE_ROWS:
            summary = f"Rows to update: {len(matched)}."
        elif plan.action == Action.DELETE_ROWS:
            summary = f"Rows to delete: {len(matched)}."
        elif plan.action == Action.CLEAR_VALUES:
            names = ", ".join(item.header for item in plan.target_columns)
            summary = f"Clear columns {names} in rows: {len(matched)}."
        else:
            raise ValueError(f"Action {plan.action.value} is not supported.")

        if is_write and len(matched) > MAX_MUTATED_ROWS:
            raise ValueError(
                f"The operation affects {len(matched)} rows; the MVP limit is "
                f"{MAX_MUTATED_ROWS}."
            )

        if plan.action in {Action.RENAME_COLUMN, Action.ADD_COLUMN, Action.DROP_COLUMN}:
            preview_rows = _preview_rows(
                worksheet,
                data_rows,
                _preview_columns(plan, sheet),
                plan,
            )
        else:
            preview_rows = _preview_rows(
                worksheet,
                matched,
                _preview_columns(plan, sheet),
                plan,
            )

        if plan.action == Action.UPDATE_ROWS:
            for row in matched:
                for assignment in plan.assignments:
                    cell = worksheet.cell(row, assignment.column.index)
                    if isinstance(cell, MergedCell):
                        raise _merged_write_error(cell, "modify")
            affected = sum(
                worksheet.cell(row, assignment.column.index).value
                != _coerce_write(
                    assignment.value,
                    worksheet.cell(row, assignment.column.index).value,
                )
                for row in matched
                for assignment in plan.assignments
            )
        elif plan.action == Action.CLEAR_VALUES:
            affected = sum(
                not _empty(worksheet.cell(row, target.index).value)
                for row in matched
                for target in plan.target_columns
            )
        elif plan.action == Action.DELETE_ROWS:
            affected = len(matched) * max(len(sheet.columns), 1)
        elif plan.action == Action.DEDUPLICATE:
            affected = len(matched) * max(len(sheet.columns), 1)
        elif plan.action == Action.DROP_COLUMN:
            target = plan.target_columns[0]
            affected = 1 + sum(
                not _empty(worksheet.cell(row, target.index).value) for row in data_rows
            )
        elif plan.action == Action.ADD_COLUMN:
            affected = 1 + (len(data_rows) if plan.fill_new_column else 0)
        elif plan.action == Action.RENAME_COLUMN:
            header_cell = worksheet.cell(sheet.header_row, plan.target_columns[0].index)
            if isinstance(header_cell, MergedCell):
                raise _merged_write_error(header_cell, "rename")
            affected = 1
        else:
            affected = 0

        has_changes = is_write and affected > 0
        return ExcelOperationPreview(
            action=plan.action,
            sheet_name=plan.sheet_name,
            is_write=is_write,
            has_changes=has_changes,
            matched_rows=len(matched),
            affected_cells=affected,
            matched_row_numbers=matched,
            rows=preview_rows,
            summary=summary,
            source_sha256=source_hash,
        )
    finally:
        workbook.close()


def _copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _merged_write_error(cell: Any, action: str) -> ValueError:
    return ValueError(
        f"Cannot {action} cell {cell.coordinate}: it belongs to a merged Excel range. "
        "Unmerge the range or target a different cell/row."
    )


def _set_regular_cell_value(cell: Any, value: Any, *, action: str, skip_merged_when_empty: bool = False) -> bool:
    """Безопасная запись с учётом openpyxl MergedCell.

    Неякорные ячейки объединённого диапазона являются read-only proxy.
    Для CLEAR их можно пропустить, если у proxy и так нет собственного значения.
    Для записей, которые должны создать/modify значение, операция блокируется
    до сохранения временной копии вместо падения AttributeError после preview.
    """
    if isinstance(cell, MergedCell):
        if skip_merged_when_empty and _empty(getattr(cell, "value", None)):
            return False
        raise _merged_write_error(cell, action)
    cell.value = value
    return True


def _apply_plan(
    worksheet: Any,
    sheet: SheetMetadata,
    plan: ExcelOperationPlan,
    preview: ExcelOperationPreview,
) -> None:
    rows = preview.matched_row_numbers
    if plan.action == Action.RENAME_COLUMN:
        target = plan.target_columns[0]
        _set_regular_cell_value(
            worksheet.cell(sheet.header_row, target.index),
            plan.new_column_name,
            action="rename",
        )
        return
    if plan.action == Action.ADD_COLUMN:
        new_index = max(column.index for column in sheet.columns) + 1
        source_index = new_index - 1
        merged_to_extend = [
            item
            for item in list(worksheet.merged_cells.ranges)
            if item.max_col == source_index
            and item.max_row < sheet.header_row
        ]
        for row_number in range(1, (worksheet.max_row or sheet.header_row) + 1):
            _copy_cell_style(
                worksheet.cell(row_number, source_index),
                worksheet.cell(row_number, new_index),
            )
        source_letter = get_column_letter(source_index)
        new_letter = get_column_letter(new_index)
        worksheet.column_dimensions[new_letter].width = (
            worksheet.column_dimensions[source_letter].width
        )
        for merged in merged_to_extend:
            worksheet.unmerge_cells(str(merged))
            worksheet.merge_cells(
                start_row=merged.min_row,
                start_column=merged.min_col,
                end_row=merged.max_row,
                end_column=new_index,
            )
        worksheet.cell(sheet.header_row, new_index).value = plan.new_column_name
        if plan.fill_new_column:
            for row_number in _data_row_numbers(worksheet, sheet):
                worksheet.cell(row_number, new_index).value = plan.new_column_default
        return
    if plan.action == Action.DROP_COLUMN:
        worksheet.delete_cols(plan.target_columns[0].index, 1)
        return
    if plan.action == Action.UPDATE_ROWS:
        for row_number in rows:
            for assignment in plan.assignments:
                cell = worksheet.cell(row_number, assignment.column.index)
                value = _coerce_write(assignment.value, getattr(cell, "value", None))
                _set_regular_cell_value(cell, value, action="modify")
        return
    if plan.action == Action.CLEAR_VALUES:
        if plan.match_cell_value is not None:
            for address in preview.matched_cell_addresses:
                cell = worksheet[address]
                _set_regular_cell_value(
                    cell,
                    None,
                    action="clear",
                    skip_merged_when_empty=True,
                )
            return
        for row_number in rows:
            for target in plan.target_columns:
                cell = worksheet.cell(row_number, target.index)
                _set_regular_cell_value(
                    cell,
                    None,
                    action="clear",
                    skip_merged_when_empty=True,
                )
        return
    if plan.action in {Action.DELETE_ROWS, Action.DEDUPLICATE}:
        if worksheet.auto_filter.ref:
            min_col, min_row, max_col, max_row = range_boundaries(
                worksheet.auto_filter.ref
            )
            deleted_inside = sum(min_row <= row <= max_row for row in rows)
            new_max_row = max(min_row, max_row - deleted_inside)
            worksheet.auto_filter.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{new_max_row}"
            )
        for row_number in sorted(rows, reverse=True):
            worksheet.delete_rows(row_number, 1)
        return
    raise ValueError(f"Cannot execute action {plan.action.value}.")


def _safe_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return {"type": "float", "value": str(value)}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return {"type": type(value).__name__, "value": str(value)}


def _workbook_digest(workbook: Any) -> str:
    payload: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        merged_ranges = list(worksheet.merged_cells.ranges)
        max_row = max([worksheet.max_row or 1, *(item.max_row for item in merged_ranges)])
        max_column = max([worksheet.max_column or 1, *(item.max_col for item in merged_ranges)])
        values = [
            [_safe_value(worksheet.cell(row, column).value) for column in range(1, max_column + 1)]
            for row in range(1, max_row + 1)
        ]
        payload.append(
            {
                "title": worksheet.title,
                "max_row": max_row,
                "max_column": max_column,
                "values": values,
                "merged": sorted(str(item) for item in merged_ranges),
                "tables": sorted(
                    (name, getattr(table, "ref", str(table)))
                    for name, table in worksheet.tables.items()
                ),
                "auto_filter": worksheet.auto_filter.ref,
            }
        )
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _verify_saved_result(
    workbook: Any,
    plan: ExcelOperationPlan,
    preview: ExcelOperationPreview,
    metadata: WorkbookMetadata,
) -> bool:
    """Проверяет смысл операции после save/reload.

    Cannot сравнивать digest книги *до* save с digest после повторного открытия:
    openpyxl при сериализации законно нормализует dimensions/metadata, поэтому
    корректный XLSX иногда выглядел как повреждённый. Здесь проверяем сам
    результат операции, а эталонный digest берём уже из сериализованной копии.
    """
    if plan.action == Action.REPLACE_ALL_VALUES:
        expected = _normalized(plan.replacement_value)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if _empty(cell.value):
                        continue
                    if _normalized(cell.value) != expected:
                        return False
        return True

    if not plan.sheet_name or plan.sheet_name not in workbook.sheetnames:
        return False
    worksheet = workbook[plan.sheet_name]
    sheet = _sheet_metadata(metadata, plan.sheet_name)

    if plan.action == Action.RENAME_COLUMN:
        target = plan.target_columns[0]
        return _normalized(worksheet.cell(sheet.header_row, target.index).value) == _normalized(plan.new_column_name)

    if plan.action == Action.ADD_COLUMN:
        headers = [worksheet.cell(sheet.header_row, col).value for col in range(1, (worksheet.max_column or 0) + 1)]
        matching = [idx for idx, value in enumerate(headers, start=1) if _normalized(value) == _normalized(plan.new_column_name)]
        if not matching:
            return False
        if plan.fill_new_column:
            col = matching[-1]
            expected = _normalized(plan.new_column_default)
            for row_number in preview.matched_row_numbers:
                if _normalized(worksheet.cell(row_number, col).value) != expected:
                    return False
        return True

    if plan.action == Action.DROP_COLUMN:
        headers = [worksheet.cell(sheet.header_row, col).value for col in range(1, (worksheet.max_column or 0) + 1)]
        target_header = plan.target_columns[0].header
        # Если в исходнике были уникальные заголовки, удалённый заголовок не должен остаться.
        original_count = sum(_normalized(column.header) == _normalized(target_header) for column in sheet.columns)
        result_count = sum(_normalized(value) == _normalized(target_header) for value in headers)
        return result_count == max(0, original_count - 1)

    if plan.action == Action.UPDATE_ROWS:
        for row_number in preview.matched_row_numbers:
            for assignment in plan.assignments:
                actual = worksheet.cell(row_number, assignment.column.index).value
                expected = _coerce_write(assignment.value, actual)
                if _normalized(actual) != _normalized(expected):
                    return False
        return True

    if plan.action == Action.CLEAR_VALUES:
        if plan.match_cell_value is not None:
            return all(
                _empty(worksheet[address].value)
                for address in preview.matched_cell_addresses
            )
        return all(
            _empty(worksheet.cell(row_number, target.index).value)
            for row_number in preview.matched_row_numbers
            for target in plan.target_columns
        )

    if plan.action in {Action.DELETE_ROWS, Action.DEDUPLICATE}:
        # Для структурных удалений точные номера строк после операции сдвигаются.
        # Проверяем ожидаемое уменьшение числа непустых data-строк, если метаданные
        # не были усечены лимитом сканирования.
        if sheet.data_rows < 10_000:
            non_empty = 0
            for row in worksheet.iter_rows(min_row=sheet.header_row + 1):
                if any(not _empty(cell.value) for cell in row):
                    non_empty += 1
            return non_empty == max(0, sheet.data_rows - preview.matched_rows)
        return True

    return True


def execute_operation(
    source_path: Path,
    original_name: str,
    operation_id: str,
    plan: ExcelOperationPlan,
    preview: ExcelOperationPreview,
    metadata: WorkbookMetadata,
    snapshots_root: Path = Path("data/snapshots/excel"),
) -> ExcelExecutionOutcome:
    if plan.action not in WRITE_ACTIONS:
        raise ValueError("A read-only operation does not require a write.")
    if _sha256(source_path) != preview.source_sha256:
        raise ValueError(
            "The source file changed after the preview. Build the plan again."
        )
    fresh_preview = build_operation_preview(source_path, plan, metadata)
    if (
        fresh_preview.matched_row_numbers != preview.matched_row_numbers
        or fresh_preview.matched_cell_addresses != preview.matched_cell_addresses
        or fresh_preview.matched_rows != preview.matched_rows
        or fresh_preview.affected_cells != preview.affected_cells
    ):
        raise ValueError("The target data changed after the preview.")

    snapshots_root.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    safe_stem = re.sub(
        r"[^0-9A-Za-zА-Яа-яЁё._() -]+",
        "_",
        Path(original_name).stem,
    ).strip(" .") or "workbook"
    snapshot_path = snapshots_root / (
        f"{safe_stem}_{timestamp}_{operation_id}_snapshot.xlsx"
    )
    temporary_path = source_path.parent / (
        f".dataops-{source_path.stem}-{operation_id}.tmp.xlsx"
    )
    restore_path = source_path.parent / (
        f".dataops-{source_path.stem}-{operation_id}.restore.xlsx"
    )
    shutil.copy2(source_path, snapshot_path)
    if _sha256(snapshot_path) != preview.source_sha256:
        raise RuntimeError("The Excel snapshot does not match the source file.")
    shutil.copy2(source_path, temporary_path)

    sheet = (
        None
        if plan.action == Action.REPLACE_ALL_VALUES
        else _sheet_metadata(metadata, plan.sheet_name or "")
    )
    columns_before = (
        sum(len(item.columns) for item in metadata.sheets)
        if sheet is None
        else len(sheet.columns)
    )
    workbook = load_workbook(temporary_path, read_only=False, data_only=False)
    try:
        if plan.action == Action.REPLACE_ALL_VALUES:
            _apply_replace_all(workbook, plan.replacement_value)
        else:
            worksheet = workbook[plan.sheet_name]
            _apply_plan(worksheet, sheet, plan, preview)
        workbook.save(temporary_path)
    finally:
        workbook.close()

    verified_workbook = load_workbook(
        temporary_path,
        read_only=False,
        data_only=False,
    )
    try:
        result_verified = _verify_saved_result(verified_workbook, plan, preview, metadata)
        # Эталон финальной проверки должен быть получен ПОСЛЕ save/reload.
        # openpyxl может нормализовать внутренние dimensions/metadata при сохранении.
        expected_digest = _workbook_digest(verified_workbook)
        if plan.action == Action.REPLACE_ALL_VALUES:
            columns_after = columns_before
        else:
            result_sheet = verified_workbook[plan.sheet_name]
            header_values = [
                result_sheet.cell(sheet.header_row, index).value
                for index in range(1, (result_sheet.max_column or 0) + 1)
            ]
            columns_after = sum(not _empty(value) for value in header_values)
    finally:
        verified_workbook.close()

    if not result_verified:
        raise RuntimeError(
            "Excel working-copy verification failed before writing to the source."
        )
    if _sha256(source_path) != preview.source_sha256:
        raise ValueError(
            "The source file changed immediately before the write. "
            "Build the plan again."
        )

    replaced = False
    try:
        os.replace(temporary_path, source_path)
        replaced = True
        final_workbook = load_workbook(source_path, read_only=False, data_only=False)
        try:
            final_verified = _workbook_digest(final_workbook) == expected_digest
        finally:
            final_workbook.close()
        if not final_verified:
            raise RuntimeError("Final Excel verification failed after the write.")
    except Exception as error:
        if replaced:
            try:
                shutil.copy2(snapshot_path, restore_path)
                os.replace(restore_path, source_path)
            except Exception as restore_error:
                raise RuntimeError(
                    "Excel was modified, verification failed, and automatic "
                    f"restore also failed: {restore_error}"
                ) from error
        raise
    finally:
        temporary_path.unlink(missing_ok=True)
        restore_path.unlink(missing_ok=True)

    return ExcelExecutionOutcome(
        snapshot_path=snapshot_path,
        result_path=source_path.resolve(),
        changed_rows=preview.matched_rows,
        affected_cells=preview.affected_cells,
        columns_before=columns_before,
        columns_after=columns_after,
        source_updated=True,
        result_verified=True,
    )
