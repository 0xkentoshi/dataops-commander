from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


output_directory = Path(__file__).parent
output_directory.mkdir(parents=True, exist_ok=True)
output_path = output_directory / "price.xlsx"

workbook = Workbook()
worksheet = workbook.active
worksheet.title = "Прайс"

worksheet.merge_cells("A1:F1")
worksheet["A1"] = "Прайс-лист магазина"
worksheet["A1"].font = Font(size=16, bold=True)
worksheet["A1"].alignment = Alignment(horizontal="center")

headers = [
    "Артикул",
    "Товар",
    "Категория",
    "Цена",
    "Остаток",
    "Обновлено",
]
for column_index, header in enumerate(headers, start=1):
    worksheet.cell(row=3, column=column_index, value=header)

header_fill = PatternFill("solid", fgColor="4F81BD")
for cell in worksheet[3]:
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = header_fill

today = date.today()
rows = [
    ["KB-001", "Клавиатура", "Периферия", 3490, 12, today],
    ["MS-002", "Мышь", "Периферия", 1890, 4, today - timedelta(days=1)],
    ["MN-003", "Монитор 24", "Мониторы", 17990, 7, today],
    ["HD-004", "Наушники", "Аудио", 5290, 0, today - timedelta(days=3)],
    ["CM-005", "Веб-камера", "Периферия", 4190, 9, today],
    ["MC-006", "Микрофон", "Аудио", 8990, 2, today - timedelta(days=2)],
]

for row in rows:
    worksheet.append(row)

for cell in worksheet["D"][3:]:
    cell.number_format = '#,##0 "₽"'

worksheet.column_dimensions["A"].width = 14
worksheet.column_dimensions["B"].width = 24
worksheet.column_dimensions["C"].width = 18
worksheet.column_dimensions["D"].width = 14
worksheet.column_dimensions["E"].width = 12
worksheet.column_dimensions["F"].width = 16

workbook.save(output_path)
print(f"Создан файл: {output_path.resolve()}")
